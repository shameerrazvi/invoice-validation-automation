"""Read and write the Excel files: the validation tool workbook and the results report."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as cfg

NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILLS = {
    "PASS": PatternFill("solid", fgColor="E2EFDA"),
    "REVIEW": PatternFill("solid", fgColor="FFF2CC"),
    "HOLD": PatternFill("solid", fgColor="F8CBAD"),
}

TEXT_COLUMNS = {
    "vendors": ["vendor_id", "vendor_name", "abn", "category", "gst_registered", "status", "bsb", "account"],
    "po_lines": ["po_number", "vendor_id", "item", "gst_code", "po_status"],
    "receipts": ["po_number"],
    "invoices": ["invoice_key", "vendor_id", "invoice_number", "abn_on_invoice", "po_number", "bsb_on_invoice",
                 "account_on_invoice"],
    "invoice_lines": ["invoice_key", "item"],
    "posted": ["vendor_id", "invoice_number"],
}
DATE_COLUMNS = {"invoices": ["invoice_date", "received_date"], "posted": ["posted_date"]}
MONEY_HINTS = ("price", "gst", "total", "line_ex", "value")

RESULT_HEADERS = ["invoice_key", "vendor_id", "vendor_name", "invoice_number", "invoice_date", "total_inc_gst",
                  "due_date", "status", "exception_codes", "details"]
EXCEPTION_HEADERS = ["invoice_key", "line_no", "code", "action", "detail"]

INSTRUCTIONS = [
    "Invoice Validation Tool",
    "",
    "Checks a day's batch of supplier invoices before they are posted for payment.",
    "",
    "How to use",
    "1. Paste the day's invoices into Invoices and Invoice_Lines (keep the column headings).",
    "2. Refresh Vendors, PO_Lines, Receipts and Posted_History from the ERP.",
    "3. Check the batch date and tolerances on the Settings sheet.",
    "4. Run the macro ValidateInvoices (Developer > Macros, or Tools > Macro on Mac).",
    "5. Work through Results: post PASS invoices, check REVIEW, resolve HOLD before posting.",
    "",
    "Status",
    "PASS    no exceptions, ready to post",
    "REVIEW  needs a person to check (second approval, old invoice) but nothing is wrong",
    "HOLD    at least one exception that must be fixed before posting",
    "",
    "Exception codes",
] + [f"{code}  {desc}" for code, desc in cfg.CODES.items()] + [
    "",
    "The same checks are available as a Python script for large batches (run_validation.py).",
]


def _style_sheet(ws, n_cols, n_rows, widths=None):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = "A2"
    if n_rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"


def _width(df, col):
    sample = [len(str(v)) for v in df[col].head(300)] if len(df) else [0]
    return min(max(len(col), *sample) + 2, 60)


def _write_df(ws, df, date_cols=(), text_cols=()):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        values = []
        for col, val in zip(df.columns, row):
            if col in date_cols and pd.notna(val) and val != "":
                val = pd.Timestamp(val).date()
            elif col in text_cols:
                val = "" if pd.isna(val) else str(val)
            elif isinstance(val, float) and pd.isna(val):
                val = None
            elif hasattr(val, "item"):
                val = val.item()
            values.append(val)
        ws.append(values)
    for c, col in enumerate(df.columns, start=1):
        letter = get_column_letter(c)
        if col in date_cols:
            for cell in ws[letter][1:]:
                cell.number_format = "dd/mm/yyyy"
        elif any(h in col for h in MONEY_HINTS) and col not in text_cols:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0.00"
    _style_sheet(ws, len(df.columns), len(df), [_width(df, c) for c in df.columns])


def build_tool_workbook(data: dict, path=cfg.TOOL_WORKBOOK) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    for r in (5, 12, 17):
        ws.cell(row=r, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 95

    st = wb.create_sheet("Settings")
    st.append(["setting", "value", "notes"])
    for key, value in cfg.SETTINGS.items():
        st.append([key, value, cfg.SETTING_NOTES[key]])
    st["B2"].number_format = "dd/mm/yyyy"
    _style_sheet(st, 3, 0, [18, 14, 70])

    for key, sheet in cfg.INPUT_SHEETS.items():
        _write_df(wb.create_sheet(sheet), data[key], DATE_COLUMNS.get(key, ()), TEXT_COLUMNS.get(key, ()))

    for sheet, headers in (("Summary", ["item", "count", "value_inc_gst"]), ("Results", RESULT_HEADERS),
                           ("Exceptions", EXCEPTION_HEADERS)):
        out = wb.create_sheet(sheet)
        out.append(headers)
        _style_sheet(out, len(headers), 0, [max(14, len(h) + 4) for h in headers])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return str(path)


def read_inputs(path=cfg.TOOL_WORKBOOK) -> tuple[dict, dict]:
    sheets = pd.read_excel(path, sheet_name=list(cfg.INPUT_SHEETS.values()) + ["Settings"],
                           dtype={c: str for cols in TEXT_COLUMNS.values() for c in cols}, keep_default_na=False)
    data = {}
    for key, sheet in cfg.INPUT_SHEETS.items():
        df = sheets[sheet]
        for col in TEXT_COLUMNS.get(key, []):
            df[col] = df[col].astype(str).replace("nan", "")
        for col in DATE_COLUMNS.get(key, []):
            df[col] = pd.to_datetime(df[col])
        data[key] = df
    settings = {}
    for _, row in sheets["Settings"].iterrows():
        val = row["value"]
        if row["setting"] == "batch_date":
            val = pd.Timestamp(val).date()
        else:
            val = float(val)
        settings[row["setting"]] = val
    return data, settings


def write_results_workbook(results, exceptions, summary, settings, path, runtime_s=None) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Invoice Validation Results"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Batch date {settings['batch_date']:%d/%m/%Y}" + (f"  |  validated in {runtime_s:.2f} s" if runtime_s else "")
    ws.append([])
    ws.append(["item", "count", "value_inc_gst"])
    for c in range(1, 4):
        ws.cell(row=4, column=c).fill, ws.cell(row=4, column=c).font = HEADER_FILL, HEADER_FONT
    for row in summary.itertuples(index=False):
        ws.append(list(row))
    for cell in ws["C"][4:]:
        cell.number_format = "#,##0.00"
    for r in range(5, 8):
        status = ws.cell(row=r, column=1).value.split()[-1]
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = STATUS_FILLS[status]
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 16

    rs = wb.create_sheet("Results")
    _write_df(rs, results, ("invoice_date", "due_date"), ("invoice_key", "vendor_id", "invoice_number"))
    status_col = RESULT_HEADERS.index("status") + 1
    for r in range(2, len(results) + 2):
        rs.cell(row=r, column=status_col).fill = STATUS_FILLS[rs.cell(row=r, column=status_col).value]
    rs.column_dimensions[get_column_letter(len(RESULT_HEADERS))].width = 100

    ex = wb.create_sheet("Exceptions")
    exc = exceptions.copy()
    exc["line_no"] = exc["line_no"].astype(object).where(exc["line_no"].notna(), None)
    _write_df(ex, exc, (), ("invoice_key", "code", "action"))
    ex.column_dimensions["E"].width = 100

    codes = wb.create_sheet("Exception Codes")
    codes.append(["code", "description", "action"])
    for code, desc in cfg.CODES.items():
        if code == "V10":
            action = "HOLD if future-dated, REVIEW if over 12 months old"
        else:
            action = "HOLD" if code in cfg.HOLD_CODES else "REVIEW"
        codes.append([code, desc, action])
    _style_sheet(codes, 3, len(cfg.CODES), [8, 60, 48])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return str(path)
